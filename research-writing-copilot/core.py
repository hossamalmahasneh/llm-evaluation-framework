"""Pure logic for Research Writing Copilot. No network calls or Streamlit dependency."""
from __future__ import annotations
import io
import json
import re
from datetime import datetime, timezone
from typing import Any
from docx import Document
from docx.shared import Inches, Pt

DEFAULT_SECTIONS = [
    'Abstract', 'Introduction', 'Related Work', 'Materials and Methods',
    'Results', 'Discussion', 'Limitations', 'Conclusion'
]
CITE_PATTERN = re.compile(r'\[@([A-Za-z0-9_.:-]+)\]')

def parse_references(raw: str) -> tuple[dict[str, str], list[str]]:
    """One source per line: KEY | bibliographic description; optional extra | fields."""
    refs: dict[str, str] = {}
    errors: list[str] = []
    for line_no, line in enumerate(raw.splitlines(), 1):
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        cols = [part.strip() for part in line.split('|')]
        if len(cols) < 2 or not cols[0] or not cols[1]:
            errors.append(f'Line {line_no}: use KEY | reference details')
            continue
        key = cols[0]
        if not re.fullmatch(r'[A-Za-z0-9_.:-]+', key):
            errors.append(f'Line {line_no}: invalid key {key!r}')
        elif key in refs:
            errors.append(f'Line {line_no}: duplicate key {key!r}')
        else:
            refs[key] = ' | '.join(cols[1:])
    return refs, errors

def citation_audit(manuscript: str, references: dict[str, str]) -> dict[str, Any]:
    keys = CITE_PATTERN.findall(manuscript)
    used = set(keys)
    known = set(references)
    return {
        'in_text_citations': len(keys),
        'distinct_citations': len(used),
        'unknown_keys': sorted(used - known),
        'unused_references': sorted(known - used),
    }

def project_template() -> dict[str, Any]:
    return {
        'title': '', 'target_journal': '', 'language': 'English', 'article_type': 'Original research',
        'research_question': '', 'contribution': '', 'method': '', 'verified_findings': '',
        'limitations': '', 'source_notes': '', 'author_voice': '', 'references': '',
        'sections': {name: '' for name in DEFAULT_SECTIONS}, 'history': [],
    }

def section_text(project: dict[str, Any]) -> str:
    return '\n\n'.join(f'## {name}\n\n{body.strip()}' for name, body in project['sections'].items() if body.strip())

def markdown_manuscript(project: dict[str, Any]) -> str:
    out = [f"# {project.get('title') or 'Untitled manuscript'}", '', section_text(project)]
    refs, _ = parse_references(project.get('references', ''))
    if refs:
        out += ['', '## References', ''] + [f'- [{key}] {detail}' for key, detail in refs.items()]
    return '\n'.join(out).strip() + '\n'

def add_revision(project: dict[str, Any], section: str, source: str, content: str) -> None:
    prev = project['sections'].get(section, '')
    project['history'].append({
        'timestamp_utc': datetime.now(timezone.utc).isoformat(), 'section': section,
        'source': source, 'previous': prev, 'new': content,
    })
    project['sections'][section] = content

def build_instructions(project: dict[str, Any], mode: str, section: str, special: str = '') -> str:
    """Generation grounded in author supplied evidence and declared bibliography."""
    refs, errors = parse_references(project.get('references', ''))
    if errors:
        raise ValueError('Fix references before generating: ' + '; '.join(errors))
    if not project.get('contribution', '').strip() and not project.get('research_question', '').strip():
        raise ValueError('Enter the research question and/or original contribution before drafting.')
    refs_block = '\n'.join(f'[@{k}] {v}' for k, v in refs.items()) or '(No references supplied)'
    prior = section_text(project)
    common = f'''You are a rigorous scientific writing collaborator. Write in {project.get('language', 'English')}.
Target journal: {project.get('target_journal', 'not specified')}; article type: {project.get('article_type', 'Original research')}.
Title: {project.get('title', '')}
Research question: {project.get('research_question', '')}
Author's original contribution: {project.get('contribution', '')}
Methodology supplied by author: {project.get('method', '')}
Author-verified results: {project.get('verified_findings', '')}
Author-declared limitations: {project.get('limitations', '')}
Source notes/grounding supplied by author: {project.get('source_notes', '')[:20000]}
Style sample written or authorized by author (use as stylistic guidance, not copied text): {project.get('author_voice', '')[:6000]}
Allowed bibliography (citation keys and descriptions):
{refs_block}
Existing manuscript sections for continuity:
{prior[:20000]}
RULES: Do not invent observations, numerical results, sample sizes, methods, sources, DOIs, or references. Cite ONLY the supplied keys in exact form [@key]; do not cite a key unless the supplied source notes support the claim. Mark unsupported factual statements as [EVIDENCE NEEDED]. Preserve substantive author contribution. Do not claim or imply that AI assistance did not occur. Output clean manuscript prose in Markdown, not meta-commentary. Do not copy the style sample. Avoid generic filler and unsupported claims. The author is responsible for reviewing every statement.'''
    if mode == 'draft':
        return common + f'\nTASK: Draft ONLY the {section} section, with a suitable heading. Detailed section instructions: {special}. Stay faithful to supplied facts.'
    if mode == 'revise':
        return common + f'\nTASK: Revise ONLY the {section} section following these explicit author editing instructions: {special}. Preserve verified facts and existing valid citation keys. Do not insert invented claims.'
    if mode == 'review':
        return common + f'\nTASK: Review the {section} section as an academic editor. Identify unsupported claims, missing evidence, weak logic, unclear methods/results, inconsistent terminology, and citation-key problems. Give prioritized, concrete comments. Do not rewrite or generate new references. Additional focus: {special}'
    raise ValueError('Unknown mode')

def export_docx(project: dict[str, Any]) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.9); sec.bottom_margin = Inches(.9)
    sec.left_margin = Inches(.9); sec.right_margin = Inches(.9)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(11)
    doc.add_heading(project.get('title', 'Untitled manuscript'), 0)
    for name, body in project.get('sections', {}).items():
        if not body.strip():
            continue
        doc.add_heading(name, level=1)
        for paragraph in body.strip().split('\n\n'):
            paragraph = re.sub(r'^#{1,6}\s+[^\n]+\n?', '', paragraph.strip()).strip()
            if paragraph:
                doc.add_paragraph(paragraph)
    refs, _ = parse_references(project.get('references', ''))
    if refs:
        doc.add_heading('References', level=1)
        for key, ref in refs.items():
            doc.add_paragraph(f'[{key}] {ref}')
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

def save_project(project: dict[str, Any]) -> str:
    return json.dumps(project, ensure_ascii=False, indent=2)

def load_project(payload: str) -> dict[str, Any]:
    data = json.loads(payload)
    if not isinstance(data, dict) or not isinstance(data.get('sections'), dict):
        raise ValueError('Invalid project JSON')
    base = project_template()
    base.update(data)
    return base
